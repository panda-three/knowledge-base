#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""课件交付验证脚本（课件共创 Skill 交付门禁的执行器）。

三种模式（可组合）：
  --notes          备注门禁：每页必须有非空演讲者备注，统计每页字数
  --spid           动画校验：p:timing 引用的 spid 必须存在于 spTree；
                   统计每页点击次数与每次点击绑定的元素数
  --render <dir>   渲染检查：LibreOffice 无头转 PDF → PyMuPDF 转 PNG，
                   输出逐页图片供人工/AI 读图核对（溢出/重叠/缺图）

用法：
  python verify_pptx.py deck.pptx --notes --spid
  python verify_pptx.py deck.pptx --render ./render_out

退出码：0 = 所选门禁全部通过；1 = 存在失败项（JSON 的 result.passed=false）。
输出：单个 JSON 对象（stdout），evidence 字段供交付前回读留档。
"""
import argparse
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def check_notes(prs):
    """备注门禁：slide 数 vs 非空备注数；返回逐页明细。"""
    pages = []
    for i, slide in enumerate(prs.slides, 1):
        text = ""
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                text = (tf.text or "").strip()
        pages.append({"page": i, "note_len": len(text), "non_empty": bool(text)})
    total = len(pages)
    non_empty = sum(1 for p in pages if p["non_empty"])
    empty_pages = [p["page"] for p in pages if not p["non_empty"]]
    return {
        "check": "notes",
        "passed": total > 0 and non_empty == total,
        "slide_count": total,
        "non_empty_note_count": non_empty,
        "empty_note_pages": empty_pages,
        "per_page": pages,
    }


def _sp_ids_in_sptree(slide):
    ids = set()
    for sp in slide.shapes:
        el = sp._element
        nv = el.find(f".//{P}cNvPr")
        if nv is None:
            nv = el.find(f".//{A}cNvPr")
        if nv is not None and nv.get("id"):
            ids.add(nv.get("id"))
    return ids


def check_spid(prs):
    """动画校验：timing 树里每个 spid 必须在 spTree 中；统计点击与编组。"""
    per_page = []
    all_ok = True
    for i, slide in enumerate(prs.slides, 1):
        sld_el = slide._element
        timing = sld_el.find(f".//{P}timing")
        sp_ids = _sp_ids_in_sptree(slide)
        if timing is None:
            per_page.append({"page": i, "has_timing": False, "clicks": 0,
                             "spids": [], "missing_spids": [], "group_sizes": []})
            continue
        xml = timing.xml if hasattr(timing, "xml") else ""
        spids = re.findall(r'spid="(\d+)"', xml)
        missing = sorted(set(s for s in spids if s not in sp_ids))
        # 每个点击（clickEffect 所在的 cTn 分支）绑定几个 spid：
        # 按 <p:par> 的 cTn childStyle/interactive 节点分组统计
        groups = []
        # p:par under p:seq/p:par/p:cTn/p:childTnLst 的每一层视为一次点击
        seqs = timing.findall(f".//{P}seq")
        click_sizes = []
        if seqs:
            # 主序列 p:seq 的 cTn/childTnLst 下每个直接 p:par 子节点 = 一次点击。
            # 注意不能用 .//{P}par 递归找——点击分支内部的嵌套 par（clickEffect/
            # withEffect 层）也带 spTgt，会把一次点击虚增成约 3 次
            for seq in seqs:
                ctn = seq.find(f"{P}cTn")
                if ctn is None:
                    continue
                ctl = ctn.find(f"{P}childTnLst")
                if ctl is None:
                    continue
                for par in ctl.findall(f"{P}par"):
                    found = par.findall(f".//{P}cBhvr")
                    sp_par = set()
                    for b in found:
                        c_spc = b.find(f".//{P}spTgt")
                        if c_spc is not None and c_spc.get("spid"):
                            sp_par.add(c_spc.get("spid"))
                    # 只统计带 spTgt 的分支（一次点击）
                    if sp_par:
                        click_sizes.append(len(sp_par))
        page_ok = not missing
        all_ok = all_ok and page_ok
        per_page.append({
            "page": i, "has_timing": True,
            "clicks": len(click_sizes),
            "spids": sorted(set(spids)),
            "missing_spids": missing,
            "group_sizes": click_sizes,
            "multi_element_clicks": [k + 1 for k, n in enumerate(click_sizes) if n > 1],
            "passed": page_ok,
        })
    multi = {p["page"]: p["multi_element_clicks"] for p in per_page if p.get("multi_element_clicks")}
    return {
        "check": "spid",
        "passed": all_ok,
        "per_page": per_page,
        "pages_with_multi_element_clicks": multi,
        "note": "multi_element_clicks 仅提示（逐元素动画页要求每次点击恰 1 个元素）；missing_spids 非空即失败",
    }


def check_render(pptx_path, out_dir):
    """渲染检查：soffice 转 PDF → PNG。返回输出目录与页数。"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pptx_path = Path(pptx_path)
    with tempfile.TemporaryDirectory() as td:
        pdf = Path(td) / (pptx_path.stem + ".pdf")
        r = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf",
             "--outdir", str(Path(td)), str(pptx_path)],
            capture_output=True, text=True, timeout=600,
        )
        # soffice 输出文件名可能与 stem 不同（含中文/空格时），找 td 里的 pdf
        pdfs = sorted(Path(td).glob("*.pdf"))
        if not pdfs:
            return {"check": "render", "passed": False,
                    "error": f"LibreOffice 未产出 PDF: rc={r.returncode} {r.stderr[:300]}"}
        pdf = pdfs[0]
        import fitz  # PyMuPDF
        doc = fitz.open(str(pdf))
        n = doc.page_count
        for idx in range(n):
            pix = doc.load_page(idx).get_pixmap(dpi=110)
            pix.save(str(out_dir / f"page_{idx + 1:02d}.png"))
        doc.close()
    return {"check": "render", "passed": True, "png_dir": str(out_dir),
            "rendered_pages": n,
            "next_step": "用读图能力逐页查看 PNG，核对溢出/重叠/空白/缺图"}


def main():
    ap = argparse.ArgumentParser(description="课件交付验证")
    ap.add_argument("pptx")
    ap.add_argument("--notes", action="store_true")
    ap.add_argument("--spid", action="store_true")
    ap.add_argument("--render", metavar="DIR", default=None)
    args = ap.parse_args()

    if not (args.notes or args.spid or args.render):
        args.notes = args.spid = True  # 默认跑两个结构门禁

    results = []
    prs = Presentation(args.pptx)
    if args.notes:
        results.append(check_notes(prs))
    if args.spid:
        results.append(check_spid(prs))
    if args.render:
        results.append(check_render(args.pptx, args.render))

    passed = all(r["passed"] for r in results)
    print(json.dumps({"file": str(Path(args.pptx).name),
                      "result": {"passed": passed},
                      "evidence": results}, ensure_ascii=False, indent=2))
    sys.exit(0 if passed else 1)


if __name__ == "__main__":
    main()
