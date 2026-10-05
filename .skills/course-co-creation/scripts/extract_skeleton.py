# -*- coding: utf-8 -*-
"""旧课件骨架提取器：逐页提取文字（按形状类型）、表格、图片、备注。
用法: python extract_skeleton.py <旧课件.pptx> [输出.txt]
输出「页序+活动位+密度」骨架表的原始材料，AI 据此整理骨架表交教师确认。
"""
import sys
from pptx import Presentation
from pptx.util import Emu


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None

    lines = []
    prs = Presentation(path)
    lines.append(f"总页数: {len(prs.slides._sldIdLst)}")
    lines.append(f"页面尺寸: {Emu(prs.slide_width).inches:.1f} x "
                 f"{Emu(prs.slide_height).inches:.1f} in")
    lines.append("=" * 70)

    for idx, slide in enumerate(prs.slides, start=1):
        lines.append(f"\n########## 第 {idx} 页 ##########")
        lines.append(f"[版式] {slide.slide_layout.name}")
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                lines.append(f"[表格 {len(tbl.rows)}行x{len(tbl.columns)}列]")
                for r in tbl.rows:
                    cells = [c.text.replace("\n", "⏎").strip() for c in r.cells]
                    lines.append("  | " + " | ".join(cells))
            elif shape.shape_type == 13:
                lines.append(f"[图片] name={shape.name} "
                             f"(内容在图片里，需读图核对，勿凭空补)")
            elif shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    is_title = shape == slide.shapes.title if slide.shapes.title else False
                    lines.append(f"[{'标题' if is_title else '文本'}] {t}")
        if slide.has_notes_slide:
            n = slide.notes_slide.notes_text_frame.text.strip()
            lines.append(f"[备注] {n if n else '(空)'}")
        else:
            lines.append("[备注] (无备注页)")

    # 密度统计
    counts = []
    for slide in prs.slides:
        c = sum(len(sh.text_frame.text.replace(" ", "").replace("\n", ""))
                for sh in slide.shapes if sh.has_text_frame)
        counts.append(c)
    lines.append("\n" + "=" * 70)
    lines.append("逐页字数: " + " ".join(f"P{i+1}={c}" for i, c in enumerate(counts)))
    lines.append(f"平均: {sum(counts)/len(counts):.1f} 字/页")

    text = "\n".join(lines)
    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"written {out_path}, {len(text)} chars")
    else:
        print(text)


if __name__ == "__main__":
    main()
