# -*- coding: utf-8 -*-
"""暗色科技风课件构建库（视觉令牌见 references/visual-tokens.md）。
用法: 在构建脚本里 `from darktech_lib import *`，逐页组装，最后 prs.save()。
动画: add_click_appear(slide, [[id], [id], ...]) —— 一组=一个元素，逐次单击出现。
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageDraw

# ---- 令牌 ----
BG     = RGBColor(0x0E, 0x0E, 0x12)
ACCENT = RGBColor(0x22, 0xD3, 0xEE)
TEXT   = RGBColor(0xE5, 0xE7, 0xEB)
SUB    = RGBColor(0x8B, 0x90, 0xA0)
CARD   = RGBColor(0x1A, 0x1C, 0x24)
BORDER = RGBColor(0x2E, 0x31, 0x40)
TERM   = RGBColor(0x0A, 0x0C, 0x10)
YH, MONO = "微软雅黑", "Consolas"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

SW, SH = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in


def new_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def make_bg(path="bg_darktech.png"):
    """深空黑 + 点阵网格 + 顶部青色微光。"""
    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), (0x0E, 0x0E, 0x12))
    d = ImageDraw.Draw(img)
    step, r = 48, 2
    for gy in range(step, H, step):
        for gx in range(step, W, step):
            d.ellipse([gx - r, gy - r, gx + r, gy + r], fill=(0x26, 0x2A, 0x36))
    for i in range(6):
        alpha = max(1, 6 - i)
        c = (int(0x22 * alpha / 8) + 10, int(0xD3 * alpha / 8) + 12,
             int(0xEE * alpha / 8) + 16)
        d.line([(0, i), (W, i)], fill=c)
    img.save(path)
    return path


def set_ea(run, name):
    """同时设拉丁+东亚字体（否则中文回退宋体）。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rPr.append(rPr.makeelement(qn('a:ea'), {'typeface': name}))


def txt(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """runs: [[(text, size, color, bold, font), ...], ...] 每个内层 list 一段。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, size, color, bold, font in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; set_ea(r, font)
    return tb


def add_bg(s, bg_path="bg_darktech.png"):
    if not os.path.exists(bg_path):
        make_bg(bg_path)
    s.shapes.add_picture(bg_path, 0, 0, width=SW, height=SH)


def header(s, tag, page):
    """页眉：青色短横线 + mono 模块标签 + 右上页码。"""
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.42),
                             Inches(0.42), Pt(3.2))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background(); bar.shadow.inherit = False
    txt(s, 1.12, 0.30, 7.5, 0.4, [[(tag, 12, ACCENT, True, MONO)]])
    txt(s, 10.4, 0.30, 2.4, 0.4, [[(page, 12, SUB, False, MONO)]],
        align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, top_light=True):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                           Inches(w), Inches(h))
    c.adjustments[0] = 0.055
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = BORDER; c.line.width = Pt(1)
    c.shadow.inherit = False
    if top_light:
        strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.28),
                                   Inches(y), Inches(w - 0.56), Pt(2.0))
        strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
        strip.line.fill.background(); strip.shadow.inherit = False
    return c


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def flow_shape(s, st, x, y, w, h, label=None, lsize=15):
    sh = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = TERM
    sh.line.color.rgb = ACCENT; sh.line.width = Pt(1.6)
    sh.shadow.inherit = False
    if label:
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.size = Pt(lsize); r.font.color.rgb = TEXT
        r.font.bold = True; set_ea(r, YH)
    return sh


def arrow(s, x1, y1, x2, y2, elbow=False):
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    ln = s.shapes.add_connector(kind, Inches(x1), Inches(y1),
                                Inches(x2), Inches(y2))
    ln.line.color.rgb = ACCENT; ln.line.width = Pt(1.8)
    ln.shadow.inherit = False
    lnEl = ln.line._get_or_add_ln()
    lnEl.append(lnEl.makeelement(qn('a:tailEnd'),
                                 {'type': 'arrow', 'w': 'med', 'len': 'med'}))
    return ln


def add_click_appear(slide, groups):
    """逐元素单击出现动画。groups: [[spid], [spid], ...]
    每组 = 一次单击，同时出现的元素放同一组——但教学节奏页一组只放一个元素。"""
    groups = [g for g in groups if g]
    if not groups:
        return
    nid = 3
    parts = []
    for spids in groups:
        sets = ""
        for spid in spids:
            sets += (f'<p:set><p:cBhvr><p:cTn id="{nid+3}" dur="1" fill="hold">'
                     f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                     f'<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                     f'</p:attrNameLst></p:cBhvr>'
                     f'<p:to><p:strVal val="visible"/></p:to></p:set>')
            nid += 1
        block = (f'<p:par><p:cTn id="{nid}" fill="hold">'
                 f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
                 f'<p:childTnLst><p:par><p:cTn id="{nid+1}" fill="hold">'
                 f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                 f'<p:childTnLst><p:par>'
                 f'<p:cTn id="{nid+2}" presetID="1" presetClass="entr" '
                 f'presetSubtype="0" fill="hold" nodeType="clickEffect">'
                 f'<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                 f'<p:childTnLst>{sets}</p:childTnLst></p:cTn></p:par>'
                 f'</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
        parts.append(block)
        nid += 3
    xml = (f'<p:timing xmlns:p="{P_NS}" xmlns:a="http://schemas.openxmlformats.'
           f'drawingml/2006/main"><p:tnLst><p:par>'
           f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
           f'<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
           f'<p:childTnLst>{"".join(parts)}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0">'
           f'<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0">'
           f'<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    slide._element.append(etree.fromstring(xml))
